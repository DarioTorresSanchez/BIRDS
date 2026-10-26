"""Trace ratios and anomalies in every Statistics tab, and Nb/U, Ce/Pb.

The per-sample tabs computed them and the column-summarising ones --
Descriptive, By group, the Paper table, the Correlation matrix -- did
not see them: a mean Eu/Eu* per group had to be worked out by hand.
They are now summarised like any column, from the same tables, and the
exported report says what the screen says.

Nb/U and Ce/Pb, the two ratios Hofmann et al. (1986) found about
constant in oceanic basalts and low in continental crust, join the
trace ratios, and so reach every tab and the data table.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
from PySide6.QtCore import Qt

from birds.normalizations import MULTIELEMENT_NORMS
from birds.plot_style import GroupStyle
from birds.reservoirs import RESERVOIRS
from birds.stats_engine import (
    RATIO_DEFINITIONS,
    RATIO_NOTES,
    compute_ratios,
    ratio_description,
)
from birds.views.statistics import (
    ALTERATION_KIND,
    ANOMALIES_KIND,
    CALCULATED_KINDS,
    RATIOS_KIND,
    calculated_columns,
)

ELEMENTS = ("La", "Ce", "Pr", "Nd", "Sm", "Eu", "Gd", "Tb", "Dy", "Ho",
            "Er", "Tm", "Yb", "Lu", "Y", "Nb", "Ta", "Th", "U", "Pb",
            "Zr", "Hf", "Ba", "Rb", "Sr", "V")
MAJORS = ("SiO2", "TiO2", "Al2O3", "FeOt", "MnO", "MgO", "CaO", "Na2O",
          "K2O", "P2O5")


@pytest.fixture
def grouped_df() -> pd.DataFrame:
    rng = np.random.default_rng(5)
    n = 12
    df = pd.DataFrame({
        "Sample_Name": [f"S{i}" for i in range(n)],
        "Group": ["mafic"] * 6 + ["felsic"] * 6,
    })
    for el in ELEMENTS + MAJORS:
        df[el] = np.abs(rng.normal(20, 5, n)) + 1
    return df


@pytest.fixture
def window(qapp, grouped_df):
    from birds.main_window import BirdsMainWindow
    w = BirdsMainWindow()
    w.data.df = grouped_df
    w.style.group_columns = ["Group"]
    w.style.groups = [GroupStyle(name=n) for n in ("mafic", "felsic")]
    w.on_dataset_replaced()
    w.statistics.refresh()
    return w


# ---------------------------------------------------------------------------
# Nb/U and Ce/Pb
# ---------------------------------------------------------------------------

def test_nb_u_and_ce_pb_are_trace_ratios():
    pairs = {name: (num, den) for name, num, den in RATIO_DEFINITIONS}
    assert pairs["Nb/U"] == ("Nb", "U")
    assert pairs["Ce/Pb"] == ("Ce", "Pb")
    df = pd.DataFrame({"Nb": [10.0, 4.0], "U": [0.2, 0.0],
                       "Ce": [50.0, 30.0], "Pb": [2.0, 3.0]})
    out = compute_ratios(df)
    assert out["Nb/U"].iloc[0] == 50.0
    assert np.isnan(out["Nb/U"].iloc[1])          # U = 0 is not a ratio
    assert list(out["Ce/Pb"]) == [25.0, 10.0]


@pytest.mark.parametrize("name, num, den, reservoir, value", [
    ("Nb/U", "Nb", "U", "N-MORB (Sun & McDonough 1989)", 50),
    ("Nb/U", "Nb", "U", "OIB (Sun & McDonough 1989)", 47),
    ("Nb/U", "Nb", "U", "Primitive mantle (McDonough & Sun 1995)", 32),
    ("Nb/U", "Nb", "U", "Bulk CC (Rudnick & Gao 2003)", 6.2),
    ("Ce/Pb", "Ce", "Pb", "N-MORB (Sun & McDonough 1989)", 25),
    ("Ce/Pb", "Ce", "Pb", "E-MORB (Sun & McDonough 1989)", 25),
    ("Ce/Pb", "Ce", "Pb", "OIB (Sun & McDonough 1989)", 25),
    ("Ce/Pb", "Ce", "Pb", "Primitive mantle (McDonough & Sun 1995)", 11),
    ("Ce/Pb", "Ce", "Pb", "Bulk CC (Rudnick & Gao 2003)", 3.9),
])
def test_the_values_in_the_tooltip_are_the_reservoirs(name, num, den,
                                                      reservoir, value):
    """Every number the description quotes is a reservoir's ratio,
    from tables checked against their papers -- not from memory."""
    r = RESERVOIRS[reservoir]
    assert r[num] / r[den] == pytest.approx(value, abs=0.5 if value > 10
                                            else 0.05)
    assert f"{value:g}" in RATIO_NOTES[name]
    assert ratio_description(name).startswith(f"{num} / {den}, ppm/ppm.")
    assert "Hofmann et al. 1986" in ratio_description(name)


def test_the_reference_is_in_the_bibliography():
    from birds.references import REFERENCES
    blob = " ".join(c for _, items in REFERENCES for c, _ in items)
    assert "10.1016/0012-821X(86)90038-5" in blob


def test_the_data_table_shows_them(window):
    window.on_view_table()
    calculated = window.data_table.model.calculated
    assert {"Nb/U", "Ce/Pb"} <= set(calculated.columns)
    expected = window.data.df["Nb"] / window.data.df["U"]
    assert np.allclose(calculated["Nb/U"], expected.round(4))


# ---------------------------------------------------------------------------
# The helper
# ---------------------------------------------------------------------------

def test_calculated_columns_leave_the_users_own_and_the_mg_number():
    df = pd.DataFrame({"La/Yb": [1.0, 2.0], "SiO2": [50.0, 60.0]})
    ratios = pd.DataFrame({"La/Yb": [9.0, 9.0], "Nb/U": [40.0, 45.0],
                           "Mg#": [60.0, 55.0], "CIA": [40.0, 42.0],
                           "Th/Nb": [np.nan, np.nan]})
    anomalies = pd.DataFrame({"Eu/Eu*": [0.8, 1.1], "Nb/U": [1.0, 1.0]})
    extra, kinds = calculated_columns(df, ratios, anomalies)
    assert list(extra.columns) == ["Nb/U", "CIA", "Eu/Eu*"]
    assert kinds == {"Nb/U": RATIOS_KIND, "CIA": ALTERATION_KIND,
                     "Eu/Eu*": ANOMALIES_KIND}
    assert list(extra["Nb/U"]) == [40.0, 45.0]     # the ratio, not a dup


# ---------------------------------------------------------------------------
# Each tab
# ---------------------------------------------------------------------------

def test_descriptive_summarises_ratios_and_anomalies(window):
    d = window.statistics.descriptive_frame
    for name in ("Nb/U", "Ce/Pb", "La/Yb", "Eu/Eu*", "(La/Yb)N", "CIA"):
        assert name in d.index, name
    ratio = window.statistics.ratios_frame["Nb/U"]
    assert d.loc["Nb/U", "mean"] == pytest.approx(ratio.mean(), abs=1e-4)


def test_by_group_summarises_them_per_group(window):
    bg = window.statistics.by_group_frame
    df = window.data.df
    nb_u = (df["Nb"] / df["U"]).round(4)
    for group in ("mafic", "felsic"):
        assert bg.loc[group, ("Nb/U", "mean")] == pytest.approx(
            nb_u[df["Group"] == group].mean(), abs=1e-4)
    assert ("Eu/Eu*", "median") in bg.columns


def test_the_paper_table_offers_them_unticked(window):
    st = window.statistics
    headers = [st.paper_tree.topLevelItem(i).text(0)
               for i in range(st.paper_tree.topLevelItemCount())]
    assert headers[2:] == list(CALCULATED_KINDS)
    items = {it.text(0): it for it in st._iter_paper_items()}
    assert items["Nb/U"].checkState(0) == Qt.CheckState.Unchecked
    items["Nb/U"].setCheckState(0, Qt.CheckState.Checked)
    items["Eu/Eu*"].setCheckState(0, Qt.CheckState.Checked)
    assert {"Nb/U", "Eu/Eu*"} <= set(st._paper.index)


def test_the_correlation_matrix_lists_them_unticked(window):
    st = window.statistics
    tops = {st.corr_tree.topLevelItem(i).text(0).split("  (")[0]:
            st.corr_tree.topLevelItem(i)
            for i in range(st.corr_tree.topLevelItemCount())}
    for kind in (RATIOS_KIND, ANOMALIES_KIND):
        assert kind in tops
        assert tops[kind].checkState(0) == Qt.CheckState.Unchecked
    assert "Nb/U" not in st.correlation_frame.columns   # not swamped
    tops[RATIOS_KIND].setCheckState(0, Qt.CheckState.Checked)
    st._update_correlation()
    assert {"Nb/U", "Ce/Pb", "La/Yb"} <= set(st.correlation_frame.columns)


def test_a_group_of_the_correlation_matrix_gets_them_too(window):
    st = window.statistics
    for i in range(st.corr_tree.topLevelItemCount()):
        st.corr_tree.topLevelItem(i).setCheckState(0, Qt.CheckState.Checked)
    st.corr_group_combo.setCurrentIndex(st.corr_group_combo.findData("mafic"))
    st._update_correlation()
    assert "Eu/Eu*" in st.correlation_frame.columns


def test_the_anomalies_follow_the_chondrite_picked(window):
    st = window.statistics
    before = st.descriptive_frame.loc["Eu/Eu*", "mean"]
    other = [st.norm_combo.itemData(i) for i in range(st.norm_combo.count())
             if st.norm_combo.itemData(i) != st.norm_combo.currentData()]
    st.norm_combo.setCurrentIndex(st.norm_combo.findData(other[0]))
    assert st.descriptive_frame.loc["Eu/Eu*", "mean"] != before
    assert st.descriptive_frame.loc["Eu/Eu*", "mean"] == pytest.approx(
        st.anomalies_frame["Eu/Eu*"].mean(), abs=1e-4)


def test_the_report_summarises_what_the_screen_does(window):
    st = window.statistics
    summary = st.summary_frame(st.norm_combo.currentData())
    assert {"Nb/U", "Ce/Pb", "Eu/Eu*"} <= set(summary.columns)
    assert not summary.columns.duplicated().any()
    # The Mg# variants are the derived ones, never the ratios table's.
    from birds.derived import with_derived
    derived = with_derived(window.data.df)
    assert ([c for c in summary.columns if c.startswith("Mg#")]
            == [c for c in derived.columns if c.startswith("Mg#")])


def test_a_reservoir_check_the_primitive_mantle_is_the_tables():
    """The Nb/U and Ce/Pb of the primitive mantle come from the same
    McDonough & Sun (1995) Table 5 values in the normaliser."""
    pm = MULTIELEMENT_NORMS["Primitive Mantle (McDonough & Sun 1995)"]
    r = RESERVOIRS["Primitive mantle (McDonough & Sun 1995)"]
    assert (pm["Nb"], pm["U"], pm["Ce"], pm["Pb"]) == (
        r["Nb"], r["U"], r["Ce"], r["Pb"])


# ---------------------------------------------------------------------------
# The Plots tab and the saved report
# ---------------------------------------------------------------------------

def test_the_plots_tab_offers_them_after_a_separator(window):
    st = window.statistics
    combo = st._plots_var_combo
    items = [combo.itemText(i) for i in range(combo.count())]
    assert {"Nb/U", "Ce/Pb", "Eu/Eu*", "CIA"} <= set(items)
    assert items.index("Nb/U") > items.index("La")
    cats = [st._plots_cat_combo.itemText(i)
            for i in range(st._plots_cat_combo.count())]
    assert "Nb/U" not in cats and "Eu/Eu*" not in cats


@pytest.mark.parametrize("kind", ["histogram", "box", "violin", "ecdf",
                                  "qq"])
@pytest.mark.parametrize("split", ["none", "overlay", "facet"])
def test_every_plot_draws_a_calculated_variable(window, kind, split):
    st = window.statistics
    st._plots_split_combo.setCurrentIndex(st._plots_split_combo.findData(split))
    st._plot_type_combo.setCurrentIndex(st._plot_type_combo.findData(kind))
    st._plots_var_combo.setCurrentIndex(st._plots_var_combo.findData("Nb/U"))
    st._render_stat_plot()
    axes = st._plots_canvas.figure.axes
    drawn = sum(len(a.patches) + len(a.lines) + len(a.collections)
                for a in axes)
    texts = " ".join(t.get_text() for a in axes for t in a.texts)
    assert drawn > 0, (kind, split)
    assert "No numeric" not in texts and "error" not in texts.lower()


def test_the_report_exports_the_matrix_on_screen(window):
    st = window.statistics
    tops = {st.corr_tree.topLevelItem(i).text(0).split("  (")[0]:
            st.corr_tree.topLevelItem(i)
            for i in range(st.corr_tree.topLevelItemCount())}
    tops[RATIOS_KIND].setCheckState(0, Qt.CheckState.Checked)
    st.corr_method_combo.setCurrentIndex(
        st.corr_method_combo.findData("spearman"))
    st._update_correlation()
    sections = st.correlation_sections()
    assert [label for label, _ in sections] == [
        "Correlation matrix (Spearman)"]
    exported = sections[0][1]
    pd.testing.assert_frame_equal(exported, st.correlation_frame)
    assert "Nb/U" in exported.columns


def test_one_sheet_per_group_when_the_screen_shows_one_per_group(window):
    st = window.statistics
    st.corr_group_combo.setCurrentIndex(
        st.corr_group_combo.findData(st._CORR_FACET))
    labels = [label for label, _ in st.correlation_sections()]
    assert labels == ["Correlation 1 - mafic", "Correlation 2 - felsic"]


def test_a_sheet_name_excel_refuses_is_cleaned(tmp_path):
    from openpyxl import load_workbook
    from birds.stats_engine import export_sections
    frame = pd.DataFrame({"a": [1.0]})
    path = tmp_path / "report.xlsx"
    export_sections(path, {"Correlation 1 - basalt: [arc]*?": frame,
                           "A/B\\C": frame}, "xlsx")
    names = load_workbook(path).sheetnames
    assert names == ["Correlation 1 - basalt_ _arc___", "A_B_C"]
