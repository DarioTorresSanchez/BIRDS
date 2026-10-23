"""The Isotopes sheet of the data template, and the ages it carries.

The Isotopes sheet now opens with the same columns, in the same order,
as the Data sheet has before its majors -- so a sample's context can be
copied across in one block -- followed by SiO2, to read the isotopes
against silica.  Both come from one list in ``docs/generate_template``.

With the context came ``Age_Unit``: the Isotopes view read every age as
Ma, so 450 ka was taken for 450 Ma in the initial-ratio correction.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from birds.dialogs.import_dialog import _merge_extra_sheet
from birds.views.isotopes import ages_in_ma

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "templates"
sys.path.insert(0, str(ROOT / "docs"))

import generate_template as gt  # noqa: E402

N_CONTEXT = len(gt.context_columns())


def _headers(sheet):
    return [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]


@pytest.fixture(scope="module")
def workbook():
    return load_workbook(TEMPLATES / "BIRDS_data_template.xlsx")


# ---------------------------------------------------------------------------
# The template
# ---------------------------------------------------------------------------

def test_both_sheets_open_with_the_same_context_columns(workbook):
    data = _headers(workbook["Data"])
    iso = _headers(workbook["Isotopes"])
    first_major = data.index("SiO2")
    assert first_major == N_CONTEXT == 32     # 28, plus the four UTM columns
    assert iso[:first_major] == data[:first_major]
    assert iso[first_major] == "SiO2"
    assert iso[first_major + 1] == "87Sr/86Sr"


def test_the_context_comes_from_one_list():
    data = [c.name for c in gt.build_columns()]
    iso = [c.name for c in gt.build_isotope_columns()]
    context = [c.name for c in gt.context_columns()]
    assert data[:N_CONTEXT] == iso[:N_CONTEXT] == context
    for name in ("Sample_Name", "Latitude", "Reference", "Age", "Age_Unit",
                 "Rock_Type", "Analytical_Method", "Notes"):
        assert name in context, name


def test_the_files_on_disk_are_what_the_generator_makes(workbook):
    iso_csv = pd.read_csv(TEMPLATES / "BIRDS_isotopes_template.csv",
                          nrows=0)
    assert list(iso_csv.columns) == [c.name for c in
                                     gt.build_isotope_columns()]
    assert _headers(workbook["Isotopes"]) == [c.name for c in
                                              gt.build_isotope_columns()]
    data_csv = pd.read_csv(TEMPLATES / "BIRDS_data_template.csv", nrows=0)
    assert list(data_csv.columns) == [c.name for c in gt.build_columns()]


def test_the_isotope_sheet_keeps_the_data_sheets_dropdowns(workbook):
    sheet = workbook["Isotopes"]
    headers = _headers(sheet)
    validated = set()
    for dv in sheet.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            validated.add(headers[rng.min_col - 1])
    assert {"Age_Unit", "Rock_Type", "TAS_Composition", "Magmatic_Series",
            "Alteration", "Analytical_Method", "Iso_Method"} <= validated


def test_the_isotope_sheet_freezes_the_same_identity_columns(workbook):
    assert workbook["Isotopes"].freeze_panes == workbook["Data"].freeze_panes


# ---------------------------------------------------------------------------
# Importing a filled template
# ---------------------------------------------------------------------------

def test_joining_the_sheets_keeps_the_data_sheets_values():
    """Repeated context columns and SiO2 come from Data; the ratios and
    a blank Group fill in from Isotopes; nothing is duplicated."""
    context = [c.name for c in gt.context_columns()]
    data = pd.DataFrame([dict.fromkeys(context, "x")] * 2)
    data["Sample_Name"] = ["A", "B"]
    data["Group"] = ["Upper", ""]
    data["SiO2"] = [52.0, 61.0]
    iso = pd.DataFrame([dict.fromkeys(context, "y")])
    iso["Sample_Name"] = ["B"]
    iso["Group"] = ["Lower"]
    iso["SiO2"] = [60.5]
    iso["87Sr/86Sr"] = [0.7051]
    merged, _note = _merge_extra_sheet(data, iso, "Isotopes")
    assert list(merged["SiO2"]) == [52.0, 61.0]
    assert list(merged["Group"]) == ["Upper", "Lower"]
    assert merged["87Sr/86Sr"].iloc[1] == pytest.approx(0.7051)
    assert not [c for c in merged.columns if c.endswith(("_extra", "_x",
                                                         "_y"))]


# ---------------------------------------------------------------------------
# Ages
# ---------------------------------------------------------------------------

def test_ages_follow_their_unit():
    df = pd.DataFrame({"Age": [450, 450, 1.2, 3e6, 12, 12, 7],
                       "Age_Unit": ["ka", "Ma", "Ga", "yr", "", None,
                                    "eons"]})
    np.testing.assert_allclose(ages_in_ma(df, "Age"),
                               [0.45, 450, 1200, 3, 12, 12, np.nan])


def test_a_column_named_in_ma_ignores_the_unit_column():
    df = pd.DataFrame({"Age_Ma": [450.0], "Age_Unit": ["ka"]})
    assert ages_in_ma(df, "Age_Ma").iloc[0] == 450.0


def test_ages_without_a_unit_column_are_ma():
    df = pd.DataFrame({"Age": ["12", "x"]})
    np.testing.assert_allclose(ages_in_ma(df, "Age"), [12, np.nan])


def test_the_isotopes_view_corrects_with_the_unit(qapp):
    from birds.data_manager import DataManager
    from birds.plot_style import PlotStyle
    from birds.views.isotopes import IsotopesView

    data = DataManager()
    data.isotopes_df = pd.DataFrame({"Sample_Name": ["A"], "Age": [450.0],
                                     "Age_Unit": ["ka"]})
    view = IsotopesView(data, PlotStyle())
    view.age_source.setCurrentIndex(view.age_source.findData("column"))
    view._populate_columns()
    view.age_column.setCurrentIndex(view.age_column.findData("Age"))
    assert view._age_series(data.isotopes_df).iloc[0] == pytest.approx(0.45)
