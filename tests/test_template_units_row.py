"""The data-entry template has no row of units, and files made with the
old one still import.

The template had the units (wt%, ppm…) in row 2, under the column names.
Read as data that row became a nameless sample, and with a handful of
samples it turned every number column to text: "wt%" among two numbers
is a third of the column.  The units are in the Reference sheet, and the
importer drops a first row that holds units and nothing else.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from birds.io_utils import read_tabular

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "BIRDS_data_template.xlsx"
sys.path.insert(0, str(ROOT / "docs"))


@pytest.mark.parametrize("sheet", ["Data", "Isotopes", "Petrography"])
def test_the_data_start_right_under_the_column_names(sheet):
    ws = load_workbook(TEMPLATE)[sheet]
    assert ws.cell(row=1, column=1).value
    assert all(c.value in (None, "") for c in ws[2])
    assert ws.freeze_panes.endswith("2")
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            assert rng.min_row == 2, (sheet, str(rng))


def test_the_units_are_in_the_reference_sheet():
    ref = load_workbook(TEMPLATE)["Reference"]
    header = [c.value for c in ref[1]]
    unit_col, name_col = header.index("Unit"), header.index("Column")
    units = {row[name_col]: row[unit_col]
             for row in ref.iter_rows(min_row=2, values_only=True)}
    assert units["SiO2"] == "wt%"
    assert units["La"] == "ppm"


def test_the_committed_template_is_the_generated_one(tmp_path, monkeypatch):
    """Regenerated from docs/generate_template.py, sheet by sheet."""
    import generate_template
    monkeypatch.setattr(generate_template, "OUT_DIR", tmp_path)
    generate_template.main()
    fresh, kept = load_workbook(tmp_path / TEMPLATE.name), load_workbook(TEMPLATE)
    for name in ("Data", "Isotopes", "Petrography"):
        a = [[c.value for c in row] for row in fresh[name].iter_rows(max_row=3)]
        b = [[c.value for c in row] for row in kept[name].iter_rows(max_row=3)]
        assert a == b, name


# ---------------------------------------------------------------------------
# Importing a file made with the old template
# ---------------------------------------------------------------------------

def _old_template_rows():
    header = ["No.", "Sample_Name", "Latitude", "SiO2", "MgO", "La", "Notes"]
    units = [None, None, "°", "wt%", "wt%", "ppm", None]
    samples = [[1, "A-1", 19.1, 50.2, 7.1, 12.0, "flow"],
               [2, "A-2", 19.2, 55.8, 4.0, 20.5, None]]
    return header, units, samples


def test_an_old_workbook_imports_without_its_units_row(tmp_path):
    header, units, samples = _old_template_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in (header, units, *samples):
        ws.append(row)
    path = tmp_path / "old.xlsx"
    wb.save(path)
    df = read_tabular(path, sheet="Data")
    assert list(df["Sample_Name"]) == ["A-1", "A-2"]
    for col in ("SiO2", "MgO", "La", "Latitude"):
        assert pd.api.types.is_numeric_dtype(df[col]), col
    assert df["SiO2"].tolist() == [50.2, 55.8]


def test_an_old_csv_imports_without_its_units_row(tmp_path):
    header, units, samples = _old_template_rows()
    lines = [",".join("" if v is None else str(v) for v in row)
             for row in (header, units, *samples)]
    path = tmp_path / "old.csv"
    path.write_text("\n".join(lines), encoding="utf-8")
    df = read_tabular(path)
    assert len(df) == 2
    assert pd.api.types.is_numeric_dtype(df["SiO2"])


@pytest.mark.parametrize("first", [
    ["A-1", "50.2", "ppm"],          # a sample, whatever else it says
    [None, "m", "°"],                # metres and degrees alone: could be data
    [None, "50.2", "7.1"],
])
def test_a_first_row_of_data_is_kept(tmp_path, first):
    path = tmp_path / "data.csv"
    rows = [["Sample_Name", "SiO2", "Notes"], first, ["A-2", "55.8", "x"]]
    path.write_text("\n".join(",".join("" if v is None else v for v in r)
                              for r in rows), encoding="utf-8")
    assert len(read_tabular(path)) == 2
