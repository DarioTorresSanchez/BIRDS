"""Render the B.I.R.D.S user manual to DOCX and PDF.

Both files share the same content (defined in ``manual_content.py``) so
they stay perfectly in sync.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
)

from manual_content import SECTIONS, SUBTITLE, TITLE


HERE = Path(__file__).resolve().parent
LOGO = HERE.parent / "birds" / "resources" / "LOGO-BIRDS.png"
FLOWCHART = HERE / "birds_flowchart.png"


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def build_docx(out_path: Path) -> None:
    doc = Document()

    # Set default font
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ---- Cover page ---------------------------------------------------
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(LOGO), width=Inches(2.6))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x3D)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(SUBTITLE)
    run.font.size = Pt(14)
    run.italic = True
    run.font.color.rgb = RGBColor(0x5B, 0x64, 0x78)

    doc.add_paragraph()
    if FLOWCHART.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(FLOWCHART), width=Inches(5.5))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run("Figure 1.  B.I.R.D.S application workflow.")
        run.font.size = Pt(9)
        run.italic = True

    doc.add_page_break()

    # ---- Body ---------------------------------------------------------
    for title_text, blocks in SECTIONS:
        h = doc.add_heading(title_text, level=1)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x3D)
        for block in blocks:
            kind = block.get("type")
            if kind == "p":
                doc.add_paragraph(block["text"])
            elif kind == "h2":
                h = doc.add_heading(block["text"], level=2)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x93)
            elif kind == "bullet":
                for item in block["items"]:
                    p = doc.add_paragraph(item, style="List Bullet")
                    p.paragraph_format.space_after = Pt(2)
            elif kind == "code":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run(block["text"])
                run.font.name = "Consolas"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x33, 0x55, 0x80)
            elif kind == "note":
                p = doc.add_paragraph()
                run = p.add_run("Note. ")
                run.bold = True
                run.font.color.rgb = RGBColor(0xB1, 0x3A, 0x3A)
                p.add_run(block["text"])
        doc.add_paragraph()  # spacer between sections

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def build_pdf(out_path: Path) -> None:
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=LETTER,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch,
        title=TITLE, author="B.I.R.D.S",
    )

    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle(
            "title", parent=base["Title"], fontSize=26, leading=32,
            textColor=colors.HexColor("#0D1B3D"), spaceAfter=8,
        ),
        "subtitle": ParagraphStyle(
            "subtitle", parent=base["Normal"], fontSize=14,
            textColor=colors.HexColor("#5B6478"), alignment=1,
            spaceAfter=20, fontName="Helvetica-Oblique",
        ),
        "h1": ParagraphStyle(
            "h1", parent=base["Heading1"], fontSize=16,
            textColor=colors.HexColor("#0D1B3D"), spaceBefore=14,
            spaceAfter=8,
        ),
        "h2": ParagraphStyle(
            "h2", parent=base["Heading2"], fontSize=12,
            textColor=colors.HexColor("#1F3A93"), spaceBefore=10,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "body", parent=base["BodyText"], fontSize=10.5, leading=14,
            spaceAfter=6, alignment=4,  # justify
        ),
        "bullet": ParagraphStyle(
            "bullet", parent=base["BodyText"], fontSize=10.5, leading=13,
            leftIndent=18, bulletIndent=6, spaceAfter=3,
        ),
        "code": ParagraphStyle(
            "code", parent=base["BodyText"], fontName="Courier",
            fontSize=9.5, leading=12, leftIndent=18,
            textColor=colors.HexColor("#33567F"),
            backColor=colors.HexColor("#F4F6FA"),
            borderColor=colors.HexColor("#D4D8E1"),
            borderWidth=0.5, borderPadding=6, spaceAfter=8,
        ),
        "note": ParagraphStyle(
            "note", parent=base["BodyText"], fontSize=10, leading=13,
            leftIndent=18, textColor=colors.HexColor("#333333"),
            backColor=colors.HexColor("#FFF7E6"),
            borderColor=colors.HexColor("#E4C170"),
            borderWidth=0.6, borderPadding=6, spaceAfter=8,
        ),
        "caption": ParagraphStyle(
            "caption", parent=base["Normal"], fontSize=9,
            alignment=1, textColor=colors.HexColor("#5B6478"),
            fontName="Helvetica-Oblique", spaceBefore=4, spaceAfter=14,
        ),
    }

    story: list = []

    # ── Cover page ────────────────────────────────────────────────────
    if LOGO.exists():
        story.append(Spacer(1, 0.6 * inch))
        story.append(Image(str(LOGO), width=2.6 * inch, height=2.6 * inch))
    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph(TITLE, styles["title"]))
    story.append(Paragraph(SUBTITLE, styles["subtitle"]))

    if FLOWCHART.exists():
        story.append(Spacer(1, 0.1 * inch))
        # Flowchart aspect ratio matches the 10.0 × 11.5 figure used by
        # generate_flowchart.py (ratio ≈ 1.15).
        story.append(Image(str(FLOWCHART),
                           width=5.6 * inch, height=6.4 * inch))
        story.append(Paragraph(
            "Figure 1.  B.I.R.D.S application workflow.",
            styles["caption"]))

    story.append(PageBreak())

    # ── Body ─────────────────────────────────────────────────────────
    for title_text, blocks in SECTIONS:
        story.append(Paragraph(_escape(title_text), styles["h1"]))
        for block in blocks:
            kind = block.get("type")
            if kind == "p":
                story.append(Paragraph(_escape(block["text"]), styles["body"]))
            elif kind == "h2":
                story.append(Paragraph(_escape(block["text"]), styles["h2"]))
            elif kind == "bullet":
                for item in block["items"]:
                    story.append(Paragraph(
                        f"• {_escape(item)}", styles["bullet"]
                    ))
            elif kind == "code":
                # Preformatted preserves whitespace + newlines
                story.append(Preformatted(block["text"], styles["code"]))
            elif kind == "note":
                story.append(Paragraph(
                    f"<b>Note.</b> {_escape(block['text'])}", styles["note"]
                ))
        story.append(Spacer(1, 0.10 * inch))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)


def _escape(text: str) -> str:
    return (text.replace("&", "&amp;")
                .replace("<", "&lt;").replace(">", "&gt;"))


def _footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#5B6478"))
    canvas.drawCentredString(
        LETTER[0] / 2.0, 0.5 * inch,
        f"B.I.R.D.S — User Manual  ·  page {doc.page}"
    )
    canvas.restoreState()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    docx_path = HERE / "BIRDS_User_Manual.docx"
    pdf_path = HERE / "BIRDS_User_Manual.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(f"DOCX: {docx_path} ({docx_path.stat().st_size} bytes)")
    print(f"PDF : {pdf_path} ({pdf_path.stat().st_size} bytes)")
